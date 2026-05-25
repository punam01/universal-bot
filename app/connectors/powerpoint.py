"""PowerPoint connector — python-pptx, one Document per slide."""
from __future__ import annotations

import io
from typing import Iterator

from pptx import Presentation

from . import connectors
from .base import Connector, Document


@connectors.register("pptx")
class PowerPointConnector(Connector):
    name = "PowerPoint (.pptx)"
    description = "Upload a PowerPoint; one Document per slide (text + speaker notes)."
    input_kind = "file"
    accepted_extensions = ["pptx"]

    def fetch(self, payload: bytes, source_name: str) -> Iterator[Document]:
        prs = Presentation(io.BytesIO(payload))
        for i, slide in enumerate(prs.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in para.runs).strip()
                        if line:
                            parts.append(line)
            # Speaker notes
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"[Speaker notes]\n{notes}")
            text = "\n".join(parts).strip()
            if text:
                yield Document(text=text, source=source_name, page=i)
