"""Word document connector — python-docx, one Document per file."""
from __future__ import annotations

import io
from typing import Iterator

from docx import Document as DocxDocument

from . import connectors
from .base import Connector, Document


@connectors.register("docx")
class WordConnector(Connector):
    name = "Word (.docx)"
    description = "Upload a Word document; paragraphs and table cells joined into one Document."
    input_kind = "file"
    accepted_extensions = ["docx"]

    def fetch(self, payload: bytes, source_name: str) -> Iterator[Document]:
        doc = DocxDocument(io.BytesIO(payload))

        parts: list[str] = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                parts.append(t)

        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    parts.append(" | ".join(row_cells))

        text = "\n\n".join(parts).strip()
        if text:
            yield Document(text=text, source=source_name)
